from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sqlite3
import sys
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


FIELD_MAP = {
    "NO": "no_value",
    "INCOMING DATE": "incoming_date",
    "DELIVERY DATE REQUESTED": "delivery_date_requested",
    "ITG-ETD": "itg_etd",
    "OVERDUE (DAYS)": "overdue_days",
    "PO NUMBER": "po_number",
    "PRODUCTCODE": "product_code",
    "CUSTOMERNAME": "customer_name",
    "ORDER REMARK": "order_remark",
    "JOBCARD": "jobcard_raw",
    "PRODUCTNAME": "product_name",
    "ORDER QTY": "order_qty",
    "PRINT_MATERIAL": "print_material",
    "PRINT_FILMWIDTH": "print_film_width",
    "PRINT_FILMTHICKNESS_ΜM": "print_film_thickness_um",
    "SLITTINGSIZE": "slitting_size",
    "PRINTINGRUBBERROLLERSIZE_1": "printing_rubber_roller_size_1",
    "PRINTINGRUBBERROLLERSIZE_2": "printing_rubber_roller_size_2",
    "PRINTINGRUBBERROLLERSIZE_3": "printing_rubber_roller_size_3",
    "PRINTINGRUBBERROLLERSIZE_4": "printing_rubber_roller_size_4",
    "COLOURQUANTITY": "colour_quantity",
    "LABELWIDTH (MM)": "label_width_mm",
    "LABELHEIGHT (MM)": "label_height_mm",
    "UPS": "ups",
    "PRINTFIlM_DENSITY".upper(): "print_film_density",
    "FINALDELIVERYFORM": "final_delivery_form",
    "ROLLLENGTHM": "roll_length_m",
    "METER RUN": "meter_run",
    "STATUS": "status",
    "PRINTING MACHINE": "printing_machine",
    "PRINTING DATE": "printing_date",
    "FILM STATUS": "film_status",
    "CYLINDER STATUS": "cylinder_status",
    "CYLINDERSUPPLIER": "cylinder_supplier",
    "CHECK CYLINDER DATE": "check_cylinder_date",
    "REMARKS": "remarks",
    "LAMINATED (PASS 1)": "laminated_pass_1",
    "LAMINATED (PASS 2)": "laminated_pass_2",
    "LAMINATED (PASS 3)": "laminated_pass_3",
    "SLITTING (SCHEDULED)": "slitting_scheduled",
    "SLITTING (ACTUAL)": "slitting_actual",
    "SEAMING (SCHEDULED)": "seaming_scheduled",
    "SEAMING (ACTUAL)": "seaming_actual",
    "INSPECTION (SCHEDULED)": "inspection_scheduled",
    "INSPECTION (ACTUAL)": "inspection_actual",
    "CUTTING (SCHEDULED)": "cutting_scheduled",
    "CUTTING (ACTUAL)": "cutting_actual",
    "PACKING (SCHEDULED)": "packing_scheduled",
    "PACKING (ACTUAL)": "packing_actual",
    "DELIVERY DATE ACTUAL": "delivery_date_actual",
    "DELIVERED QTY": "delivered_qty",
    "DO NUM": "do_number",
    "STOCK BALANCE": "stock_balance",
    "COLUMN1": "extra_column_1",
    "COLUMN2": "extra_column_2",
    "KG/MTR": "kg_per_meter",
}

JOB_FIELDS = [
    "jobcard_raw", "jobcard_normalized", "jobcard_last4", "no_value",
    "incoming_date", "delivery_date_requested", "itg_etd", "overdue_days",
    "po_number", "product_code", "customer_name", "order_remark", "product_name",
    "order_qty", "print_material", "print_film_width", "print_film_thickness_um",
    "slitting_size", "printing_rubber_roller_size_1", "printing_rubber_roller_size_2",
    "printing_rubber_roller_size_3", "printing_rubber_roller_size_4", "colour_quantity",
    "label_width_mm", "label_height_mm", "ups", "print_film_density",
    "final_delivery_form", "roll_length_m", "meter_run", "status", "printing_machine",
    "printing_date", "film_status", "cylinder_status", "cylinder_supplier",
    "check_cylinder_date", "remarks", "laminated_pass_1", "laminated_pass_2",
    "laminated_pass_3", "slitting_scheduled", "slitting_actual", "seaming_scheduled",
    "seaming_actual", "inspection_scheduled", "inspection_actual", "cutting_scheduled",
    "cutting_actual", "packing_scheduled", "packing_actual", "delivery_date_actual",
    "delivered_qty", "do_number", "stock_balance", "extra_column_1", "extra_column_2",
    "kg_per_meter",
]

REQUIRED_HEADERS = {"JOBCARD", "PRODUCTCODE", "CUSTOMERNAME", "PRODUCTNAME", "METER RUN"}


def normalized_header(value: Any) -> str:
    if value is None:
        return ""
    original = str(value)
    text = re.sub(r"\s+", " ", original.replace("\r", " ").replace("\n", " ")).strip().upper()
    if text == "DELIVERY DATE":
        return "DELIVERY DATE REQUESTED" if original[-1:].isspace() else "DELIVERY DATE ACTUAL"
    return text


def serializable(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def text_value(value: Any) -> str | None:
    value = serializable(value)
    if value is None:
        return None
    if isinstance(value, bool):
        return "1" if value else "0"
    return str(value).strip()


def normalize_jobcard(value: str) -> str | None:
    normalized = re.sub(r"\s+", " ", value.strip().upper())
    match = re.search(r"(?:ITG|SA)[- ]?\d{4}[- ]?\d{4}", normalized)
    if match:
        return re.sub(r"\s+", "", match.group(0))
    compact = re.sub(r"[^A-Z0-9-]", "", normalized)
    return compact or None


def last_four(value: str | None) -> str | None:
    if not value:
        return None
    values = re.findall(r"\d{4}", value)
    return values[-1] if values else None


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def stable_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def resolve_path(config_path: Path, value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else (config_path.parent / path).resolve()


def read_snapshot(source: Path, sheet_name: str, header_row: int) -> dict[str, Any]:
    if not source.is_file():
        raise FileNotFoundError(source)
    if "2026B" not in source.name.upper():
        raise ValueError(f"Expected the live 2026B workbook, received: {source.name}")

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet {sheet_name!r} not found in {source.name}")
        worksheet = workbook[sheet_name]
        raw_headers = list(next(worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True)))
        last_header_index = max((i for i, value in enumerate(raw_headers) if value not in (None, "")), default=-1)
        raw_headers = raw_headers[: last_header_index + 1]
        headers = [normalized_header(value) for value in raw_headers]
        missing = sorted(REQUIRED_HEADERS.difference(headers))
        if missing:
            raise ValueError(f"Required headers missing: {', '.join(missing)}")

        valid_rows = []
        rejected_rows = []
        ignored_blank = 0
        for source_row, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, max_col=len(headers), values_only=True),
            start=header_row + 1,
        ):
            values = [serializable(value) for value in row]
            if not any(value not in (None, "") for value in values):
                ignored_blank += 1
                continue
            payload_json = stable_json({
                "headers": [serializable(value) for value in raw_headers],
                "values": {f"column_{i + 1:03d}": value for i, value in enumerate(values)},
            })
            mapped = {field: None for field in JOB_FIELDS}
            for header, value in zip(headers, row):
                destination = FIELD_MAP.get(header)
                if destination:
                    mapped[destination] = text_value(value)
            jobcard = mapped.get("jobcard_raw")
            record = {"source_row": source_row, "payload_json": payload_json, "mapped": mapped}
            if jobcard is None or not jobcard.strip():
                record["rule_code"] = "MISSING_JOBCARD"
                record["reason"] = "The source row contains data but JOBCARD is blank."
                rejected_rows.append(record)
                continue
            mapped["jobcard_normalized"] = normalize_jobcard(jobcard)
            mapped["jobcard_last4"] = last_four(mapped["jobcard_normalized"])
            valid_rows.append(record)
        return {
            "headers": raw_headers,
            "valid_rows": valid_rows,
            "rejected_rows": rejected_rows,
            "ignored_blank": ignored_blank,
        }
    finally:
        workbook.close()


def validate_snapshot(snapshot: dict[str, Any], previous_count: int, minimum: int, max_drop: float) -> None:
    current = len(snapshot["valid_rows"])
    if current < minimum:
        raise ValueError(f"Only {current} valid jobs found; configured minimum is {minimum}")
    if previous_count > 0:
        drop_percent = ((previous_count - current) / previous_count) * 100
        if drop_percent > max_drop:
            raise ValueError(
                f"Valid-job count fell from {previous_count} to {current} "
                f"({drop_percent:.1f}%); configured maximum drop is {max_drop:.1f}%"
            )


def insert_source_row(
    connection: sqlite3.Connection,
    batch_id: int,
    source_name: str,
    source_hash: str,
    sheet_name: str,
    record: dict[str, Any],
) -> int:
    payload = record["payload_json"]
    row_hash = hashlib.sha256(payload.encode("utf-8")).hexdigest()
    cursor = connection.execute(
        """
        INSERT INTO live_source_rows (
            import_batch_id, source_file, source_file_sha256, source_sheet,
            source_row, source_row_sha256, raw_payload_json
        ) VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (batch_id, source_name, source_hash, sheet_name, record["source_row"], row_hash, payload),
    )
    return int(cursor.lastrowid)


def write_reports(connection: sqlite3.Connection, report_dir: Path, summary: dict[str, Any]) -> None:
    report_dir.mkdir(parents=True, exist_ok=True)
    (report_dir / "live_import_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    rows = connection.execute(
        """
        SELECT s.source_file, s.source_sheet, s.source_row, r.rule_code, r.reason
        FROM live_import_rejections AS r
        JOIN live_source_rows AS s ON s.id = r.source_row_id
        ORDER BY s.source_row
        """
    ).fetchall()
    with (report_dir / "live_import_rejections.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["source_file", "source_sheet", "source_row", "rule_code", "reason"])
        writer.writerows(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description="Refresh ITG live schedule SQLite database.")
    parser.add_argument("--config", required=True)
    parser.add_argument("--source", help="Optional source XLSX override")
    parser.add_argument("--database", help="Optional SQLite output override")
    args = parser.parse_args()

    config_path = Path(args.config).resolve()
    config = json.loads(config_path.read_text(encoding="utf-8"))
    source = Path(args.source).resolve() if args.source else resolve_path(config_path, config["source"])
    database = Path(args.database).resolve() if args.database else resolve_path(config_path, config["database"])
    reports = resolve_path(config_path, config["reports_directory"])
    sheet_name = str(config["sheet"])
    header_row = int(config.get("header_row", 2))
    minimum = int(config.get("minimum_valid_rows", 1))
    maximum_drop = float(config.get("maximum_drop_percent", 50))

    snapshot = read_snapshot(source, sheet_name, header_row)
    source_hash = file_sha256(source)
    source_size = source.stat().st_size
    database.parent.mkdir(parents=True, exist_ok=True)
    schema_path = Path(__file__).resolve().parents[1] / "sql" / "schema.sql"
    connection = sqlite3.connect(database)
    connection.execute("PRAGMA foreign_keys = ON")
    try:
        connection.executescript(schema_path.read_text(encoding="utf-8"))
        previous_count = int(connection.execute("SELECT COUNT(*) FROM live_jobs").fetchone()[0])
        validate_snapshot(snapshot, previous_count, minimum, maximum_drop)
        started = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
        batch_key = hashlib.sha256(f"{source_hash}\n{sheet_name}".encode("utf-8")).hexdigest()

        connection.execute("BEGIN IMMEDIATE")
        connection.execute(
            """
            INSERT INTO live_import_batches (
                batch_key, source_file, source_file_sha256, source_sheet,
                source_file_size_bytes, started_at_utc, status
            ) VALUES (?, ?, ?, ?, ?, ?, 'running')
            ON CONFLICT(batch_key) DO UPDATE SET
                source_file=excluded.source_file,
                source_file_sha256=excluded.source_file_sha256,
                source_sheet=excluded.source_sheet,
                source_file_size_bytes=excluded.source_file_size_bytes,
                started_at_utc=excluded.started_at_utc,
                completed_at_utc=NULL,
                status='running',
                rows_seen=0, rows_loaded=0, rows_rejected=0, rows_ignored_blank=0
            """,
            (batch_key, source.name, source_hash, sheet_name, source_size, started),
        )
        batch_id = int(connection.execute(
            "SELECT id FROM live_import_batches WHERE batch_key=?", (batch_key,)
        ).fetchone()[0])

        connection.execute("DELETE FROM live_source_rows")
        columns = ["source_row_id", *JOB_FIELDS]
        placeholders = ", ".join("?" for _ in columns)
        job_sql = f"INSERT INTO live_jobs ({', '.join(columns)}) VALUES ({placeholders})"

        for record in snapshot["valid_rows"]:
            source_row_id = insert_source_row(
                connection, batch_id, source.name, source_hash, sheet_name, record
            )
            mapped = record["mapped"]
            connection.execute(job_sql, [source_row_id, *(mapped.get(field) for field in JOB_FIELDS)])

        for record in snapshot["rejected_rows"]:
            source_row_id = insert_source_row(
                connection, batch_id, source.name, source_hash, sheet_name, record
            )
            connection.execute(
                "INSERT INTO live_import_rejections (source_row_id, rule_code, reason) VALUES (?, ?, ?)",
                (source_row_id, record["rule_code"], record["reason"]),
            )

        completed = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
        seen = len(snapshot["valid_rows"]) + len(snapshot["rejected_rows"])
        connection.execute(
            """
            UPDATE live_import_batches
            SET completed_at_utc=?, status='completed', rows_seen=?, rows_loaded=?,
                rows_rejected=?, rows_ignored_blank=?
            WHERE id=?
            """,
            (
                completed, seen, len(snapshot["valid_rows"]), len(snapshot["rejected_rows"]),
                snapshot["ignored_blank"], batch_id,
            ),
        )
        connection.commit()

        summary = {
            "batch_key": batch_key,
            "completed_at_utc": completed,
            "source_file": source.name,
            "source_file_sha256": source_hash,
            "source_sheet": sheet_name,
            "previous_live_jobs": previous_count,
            "rows_seen": seen,
            "rows_loaded": len(snapshot["valid_rows"]),
            "rows_rejected": len(snapshot["rejected_rows"]),
            "rows_ignored_blank": snapshot["ignored_blank"],
            "database": database.name,
        }
        write_reports(connection, reports, summary)
        print(json.dumps(summary, indent=2, sort_keys=True))
        return 0
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)

