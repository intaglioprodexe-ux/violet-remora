#!/usr/bin/env python3
"""Import tbl_SourceMasterlist from an XLSM workbook into SQLite safely."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import sqlite3
import sys
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries
except ImportError as exc:
    raise SystemExit(
        "openpyxl is required. Run: py -m pip install -r requirements.txt"
    ) from exc


REQUIRED_COLUMNS = {
    "ProductStatus",
    "DateCreated",
    "SuppliedBy",
    "RouteID",
    "Category",
    "APPLICATION",
    "ProductLayers",
    "ProductCategory",
    "CustomerName",
    "ProductName",
    "ProductCode",
    "LabelWidth (MM)",
    "LabelHeight (MM)",
    "UPS",
    "ColourQuantity",
    "SlittingSize",
    "RollLengthM",
    "Print_Material",
    "Print_FilmWidth",
    "Print_FilmThickness_µm",
    "PrintFIlm_Density",
    "Barrier1_FilmWidth",
    "Barrier1_Thickness_µm",
    "Barrier1_Density",
    "Barrier2_FIlmWidth",
    "Barrier2_Thickness_µm",
    "Barrier2_Density",
    "Sealant_FilmWidth",
    "Sealant_Thickness_µm",
    "Sealant_Density",
}

ITEM_COLUMNS = (
    "import_batch_id", "source_row_id", "source_excel_row", "product_code_raw",
    "product_code", "product_code_normalized", "product_status", "date_created",
    "supplied_by", "route_id", "category", "application", "product_layers",
    "product_category", "customer_name", "product_name", "label_width_mm",
    "label_height_mm", "ups", "colour_quant_raw", "colour_quantity",
    "slitting_size", "roll_length_m", "print_material", "print_film_width",
    "print_film_thickness_um_raw", "print_film_thickness_um", "print_film_density",
    "barrier1_film_width", "barrier1_thickness_um", "barrier1_density",
    "barrier2_film_width", "barrier2_thickness_um", "barrier2_density",
    "sealant_film_width", "sealant_thickness_um", "sealant_density",
    "l1p1t1", "l2p2t2", "l3p3t3", "l4p4t4", "lpt", "m2r", "m2p",
    "l1p1t1w1", "l2p2t2w2", "l3p3t3w3", "l4p4t4w4", "pcs_kg", "m_kg",
    "jm_kg", "pt_kg", "process_51_printing", "process_52_laminating",
    "process_53_slitting", "process_54_seaming", "process_55_inspecting",
    "process_56_cutting", "process_57_packing", "calculation_warnings_json",
    "raw_payload_json", "imported_at",
)


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    return value


def stable_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def text_or_none(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def number_or_none(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        result = float(value)
        return result if math.isfinite(result) else None
    text = str(value).strip().replace(",", "")
    if not text or text == "-":
        return None
    try:
        result = float(text)
        return result if math.isfinite(result) else None
    except ValueError:
        return None


def int_or_none(value: Any) -> int | None:
    number = number_or_none(value)
    if number is None:
        return None
    return int(number)


def layer_weight(thickness: Any, density: Any) -> float:
    thickness_number = number_or_none(thickness)
    density_number = number_or_none(density)
    if thickness_number is None or density_number is None:
        return 0.0
    return (thickness_number / 1000.0) * (density_number / 1_000_000.0)


def multiply_or_none(*values: Any) -> float | None:
    result = 1.0
    for value in values:
        number = number_or_none(value)
        if number is None:
            return None
        result *= number
    return result


def zero_if_none(value: float | None) -> float:
    return 0.0 if value is None else value


def colour_count(value: Any) -> int:
    text = text_or_none(value) or "0"
    try:
        return int(text[0])
    except (ValueError, IndexError):
        return 0


def calculate_item(raw: dict[str, Any]) -> dict[str, Any]:
    warnings: list[str] = []
    route = text_or_none(raw.get("RouteID")) or ""
    category = text_or_none(raw.get("ProductCategory")) or ""
    material = text_or_none(raw.get("Print_Material"))
    thickness_raw = number_or_none(raw.get("Print_FilmThickness_µm"))
    thickness = thickness_raw
    if material in {"PVC BLOW(54%)", "PVC - BLOW"} and thickness_raw == 40:
        thickness = 38.0

    l1 = layer_weight(thickness, raw.get("PrintFIlm_Density"))
    l2 = layer_weight(raw.get("Barrier1_Thickness_µm"), raw.get("Barrier1_Density"))
    l3 = layer_weight(raw.get("Barrier2_Thickness_µm"), raw.get("Barrier2_Density"))
    l4 = layer_weight(raw.get("Sealant_Thickness_µm"), raw.get("Sealant_Density"))
    lpt = l1 + l2 + l3 + l4

    slit = number_or_none(raw.get("SlittingSize"))
    roll_length = number_or_none(raw.get("RollLengthM"))
    width = number_or_none(raw.get("LabelWidth (MM)"))
    height = number_or_none(raw.get("LabelHeight (MM)"))
    ups = number_or_none(raw.get("UPS"))

    m2r = None if slit is None or roll_length is None else slit * (roll_length * 1000.0)
    dimension = width if route in {"LAMINATION", "NON SHRINK"} else height
    m2p = None if slit is None or dimension is None else dimension * slit
    if m2r is None:
        warnings.append("M2R unavailable: SlittingSize or RollLengthM is not numeric")
    if m2p is None:
        warnings.append("M2P unavailable: SlittingSize or route dimension is not numeric")

    l1w1 = multiply_or_none(thickness, 1 / 1000.0, raw.get("Print_FilmWidth"), raw.get("PrintFIlm_Density"), 1 / 1_000_000.0)
    l2w2 = multiply_or_none(raw.get("Barrier1_FilmWidth"), raw.get("Barrier1_Thickness_µm"), 1 / 1000.0, raw.get("Barrier1_Density"), 1 / 1_000_000.0)
    l3w3 = multiply_or_none(raw.get("Barrier2_Thickness_µm"), 1 / 1000.0, raw.get("Barrier2_Density"), 1 / 1_000_000.0, raw.get("Barrier2_FIlmWidth"))
    l4w4 = multiply_or_none(raw.get("Sealant_FilmWidth"), raw.get("Sealant_Thickness_µm"), 1 / 1000.0, raw.get("Sealant_Density"), 1 / 1_000_000.0)

    pcs_dimension = width if route in {"NON SHRINK", "LAMINATION"} else height
    pcs_kg = zero_if_none(multiply_or_none(lpt, slit, pcs_dimension))
    m_kg = zero_if_none(multiply_or_none(lpt, slit, 1000.0))
    jm_width = None if slit is None or ups is None else (slit * ups) + 20.0
    jm_kg = zero_if_none(multiply_or_none(lpt, jm_width, 1000.0))
    pt_kg = zero_if_none(multiply_or_none(l1, jm_width, 1000.0))

    return {
        "label_width_mm": width,
        "label_height_mm": height,
        "ups": int_or_none(raw.get("UPS")),
        "colour_quant_raw": text_or_none(raw.get("ColourQuantity")),
        "colour_quantity": colour_count(raw.get("ColourQuantity")),
        "slitting_size": slit,
        "roll_length_m": roll_length,
        "print_material": material,
        "print_film_width": number_or_none(raw.get("Print_FilmWidth")),
        "print_film_thickness_um_raw": thickness_raw,
        "print_film_thickness_um": thickness,
        "print_film_density": number_or_none(raw.get("PrintFIlm_Density")),
        "barrier1_film_width": number_or_none(raw.get("Barrier1_FilmWidth")),
        "barrier1_thickness_um": number_or_none(raw.get("Barrier1_Thickness_µm")),
        "barrier1_density": number_or_none(raw.get("Barrier1_Density")),
        "barrier2_film_width": number_or_none(raw.get("Barrier2_FIlmWidth")),
        "barrier2_thickness_um": number_or_none(raw.get("Barrier2_Thickness_µm")),
        "barrier2_density": number_or_none(raw.get("Barrier2_Density")),
        "sealant_film_width": number_or_none(raw.get("Sealant_FilmWidth")),
        "sealant_thickness_um": number_or_none(raw.get("Sealant_Thickness_µm")),
        "sealant_density": number_or_none(raw.get("Sealant_Density")),
        "l1p1t1": l1, "l2p2t2": l2, "l3p3t3": l3, "l4p4t4": l4,
        "lpt": lpt, "m2r": m2r, "m2p": m2p,
        "l1p1t1w1": l1w1, "l2p2t2w2": l2w2, "l3p3t3w3": l3w3,
        "l4p4t4w4": l4w4, "pcs_kg": pcs_kg, "m_kg": m_kg,
        "jm_kg": jm_kg, "pt_kg": pt_kg,
        "process_51_printing": 1 if category == "PRINTED" else 0,
        "process_52_laminating": 1 if route == "LAMINATION" else 0,
        "process_53_slitting": 1,
        "process_54_seaming": 1 if route in {"SHRINK PIECE", "SHRINK ROLL"} else 0,
        "process_55_inspecting": 1 if route == "SHRINK ROLL" else 0,
        "process_56_cutting": 1 if route == "SHRINK PIECE" else 0,
        "process_57_packing": 1,
        "calculation_warnings_json": stable_json(warnings),
    }


def load_config(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        config = json.load(handle)
    base = path.resolve().parent
    for key in ("source_file", "database_path"):
        config[key] = os.path.expandvars(config[key])
    for key in ("summary_report", "rejection_report"):
        candidate = Path(os.path.expandvars(config[key]))
        config[key] = str(candidate if candidate.is_absolute() else base / candidate)
    return config


def open_table(source: Path, worksheet: str, table_name: str):
    formula_book = load_workbook(source, data_only=False, read_only=False, keep_vba=True, keep_links=False)
    value_book = load_workbook(source, data_only=True, read_only=False, keep_vba=True, keep_links=False)
    if worksheet not in formula_book.sheetnames:
        raise ValueError(f"Worksheet not found: {worksheet}")
    formula_sheet = formula_book[worksheet]
    value_sheet = value_book[worksheet]
    if table_name not in formula_sheet.tables:
        raise ValueError(f"Excel table not found: {table_name}")
    table = formula_sheet.tables[table_name]
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    headers = [formula_sheet.cell(min_row, column).value for column in range(min_col, max_col + 1)]
    if len(headers) != len(set(headers)):
        raise ValueError("The Excel table contains duplicate column headings")
    missing = sorted(REQUIRED_COLUMNS - set(headers))
    if missing:
        raise ValueError("Required columns are missing: " + ", ".join(missing))
    return formula_book, value_book, formula_sheet, value_sheet, headers, min_col, min_row, max_col, max_row


def import_masterlist(config: dict[str, Any], source_override: str | None, database_override: str | None) -> dict[str, Any]:
    source = Path(os.path.expandvars(source_override or config["source_file"])).resolve()
    database = Path(os.path.expandvars(database_override or config["database_path"])).resolve()
    if not source.is_file():
        raise FileNotFoundError(f"Source workbook not found: {source}")
    database.parent.mkdir(parents=True, exist_ok=True)
    schema_path = Path(__file__).resolve().parents[1] / "sql" / "schema.sql"
    source_hash = sha256_file(source)
    imported_at = utc_now()
    batch_key = f"{source_hash}:{config['worksheet']}:{config['table']}"

    books = open_table(source, config["worksheet"], config["table"])
    formula_book, value_book, formula_sheet, value_sheet, headers, min_col, min_row, max_col, max_row = books
    rows_seen = max_row - min_row
    if rows_seen < int(config.get("minimum_source_rows", 1)):
        raise ValueError(f"Source validation failed: only {rows_seen} rows; minimum is {config['minimum_source_rows']}")

    connection = sqlite3.connect(database)
    connection.row_factory = sqlite3.Row
    try:
        connection.executescript(schema_path.read_text(encoding="utf-8"))
        existing = connection.execute(
            "SELECT id, imported_at, rows_loaded, rows_rejected, rows_deduplicated FROM import_batches WHERE batch_key = ? AND status = 'success'",
            (batch_key,),
        ).fetchone()
        if existing:
            existing_rejections = [
                dict(row) for row in connection.execute(
                    """SELECT source_row, product_code_raw, severity, rule_code, reason
                    FROM import_rejections WHERE import_batch_id = ? ORDER BY source_row, id""",
                    (existing["id"],),
                )
            ]
            return {
                "status": "skipped_unchanged",
                "database": str(database), "source_file": str(source), "source_hash": source_hash,
                "batch_id": existing["id"], "imported_at": existing["imported_at"],
                "rows_seen": rows_seen, "rows_loaded": existing["rows_loaded"],
                "rows_rejected": existing["rows_rejected"],
                "rows_deduplicated": existing["rows_deduplicated"],
                "rejections": existing_rejections,
                "message": "The workbook hash was already imported; current data was left unchanged.",
            }

        connection.execute("BEGIN IMMEDIATE")
        cursor = connection.execute(
            """INSERT INTO import_batches (
                batch_key, source_file, source_hash, source_size_bytes, source_modified_at,
                imported_at, source_sheet, source_table, status
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'running')""",
            (batch_key, str(source), source_hash, source.stat().st_size,
             datetime.fromtimestamp(source.stat().st_mtime, timezone.utc).replace(microsecond=0).isoformat(),
             imported_at, config["worksheet"], config["table"]),
        )
        batch_id = cursor.lastrowid
        connection.execute("DELETE FROM master_items")
        seen_codes: dict[str, int] = {}
        loaded = rejected = deduplicated = 0
        rejection_rows: list[dict[str, Any]] = []
        placeholders = ",".join("?" for _ in ITEM_COLUMNS)
        item_sql = f"INSERT INTO master_items ({','.join(ITEM_COLUMNS)}) VALUES ({placeholders})"

        for excel_row in range(min_row + 1, max_row + 1):
            raw: dict[str, Any] = {}
            formulas: dict[str, str] = {}
            for offset, header in enumerate(headers):
                value_cell = value_sheet.cell(excel_row, min_col + offset)
                formula_cell = formula_sheet.cell(excel_row, min_col + offset)
                raw[str(header)] = json_value(value_cell.value)
                if formula_cell.data_type == "f":
                    formulas[str(header)] = str(formula_cell.value)
            raw_json = stable_json(raw)
            formula_json = stable_json(formulas)
            row_hash = hashlib.sha256(
                stable_json({"sheet": config["worksheet"], "table": config["table"], "row": excel_row, "values": raw}).encode("utf-8")
            ).hexdigest()
            source_cursor = connection.execute(
                """INSERT INTO source_rows (
                    import_batch_id, source_sheet, source_table, source_row, source_row_hash,
                    raw_payload_json, formula_payload_json
                ) VALUES (?, ?, ?, ?, ?, ?, ?)""",
                (batch_id, config["worksheet"], config["table"], excel_row, row_hash, raw_json, formula_json),
            )
            source_row_id = source_cursor.lastrowid
            product_code_raw_value = raw.get("ProductCode")
            product_code = text_or_none(product_code_raw_value)
            product_code_raw = "" if product_code_raw_value is None else str(product_code_raw_value)
            normalized = product_code.upper() if product_code else ""
            if not normalized:
                reason = "ProductCode is blank after text normalization"
                rule = "MISSING_PRODUCT_CODE"
            elif normalized in seen_codes:
                reason = f"Duplicate ProductCode; first occurrence is Excel row {seen_codes[normalized]}"
                rule = "DUPLICATE_PRODUCT_CODE"
            else:
                reason = rule = ""

            if rule:
                connection.execute(
                    """INSERT INTO import_rejections (
                        import_batch_id, source_row_id, source_row, product_code_raw,
                        severity, rule_code, reason, raw_payload_json
                    ) VALUES (?, ?, ?, ?, 'rejected', ?, ?, ?)""",
                    (batch_id, source_row_id, excel_row, product_code_raw, rule, reason, raw_json),
                )
                rejected += 1
                deduplicated += rule == "DUPLICATE_PRODUCT_CODE"
                rejection_rows.append({
                    "source_row": excel_row, "product_code_raw": product_code_raw,
                    "severity": "rejected", "rule_code": rule, "reason": reason,
                })
                continue

            seen_codes[normalized] = excel_row
            calculated = calculate_item(raw)
            item = {
                "import_batch_id": batch_id, "source_row_id": source_row_id,
                "source_excel_row": excel_row, "product_code_raw": product_code_raw,
                "product_code": product_code, "product_code_normalized": normalized,
                "product_status": text_or_none(raw.get("ProductStatus")),
                "date_created": json_value(raw.get("DateCreated")),
                "supplied_by": text_or_none(raw.get("SuppliedBy")),
                "route_id": text_or_none(raw.get("RouteID")),
                "category": text_or_none(raw.get("Category")),
                "application": text_or_none(raw.get("APPLICATION")),
                "product_layers": text_or_none(raw.get("ProductLayers")),
                "product_category": text_or_none(raw.get("ProductCategory")),
                "customer_name": text_or_none(raw.get("CustomerName")),
                "product_name": text_or_none(raw.get("ProductName")),
                "raw_payload_json": raw_json, "imported_at": imported_at,
                **calculated,
            }
            connection.execute(item_sql, tuple(item[column] for column in ITEM_COLUMNS))
            loaded += 1

        minimum_ratio = float(config.get("minimum_loaded_ratio", 0.0))
        if rows_seen == 0 or loaded / rows_seen < minimum_ratio:
            raise ValueError(
                f"Load validation failed: {loaded}/{rows_seen} rows loaded; required ratio is {minimum_ratio:.2f}"
            )
        connection.execute(
            """UPDATE import_batches SET rows_seen = ?, rows_loaded = ?, rows_rejected = ?,
                rows_deduplicated = ?, status = 'success', message = ? WHERE id = ?""",
            (rows_seen, loaded, rejected, deduplicated, "Current master list replaced transactionally", batch_id),
        )
        retention = max(1, int(config.get("retention_batches", 7)))
        connection.execute(
            """DELETE FROM import_batches WHERE id IN (
                SELECT id FROM import_batches WHERE status = 'success' ORDER BY imported_at DESC, id DESC LIMIT -1 OFFSET ?
            )""",
            (retention,),
        )
        connection.commit()
        return {
            "status": "success", "database": str(database), "source_file": str(source),
            "source_hash": source_hash, "batch_id": batch_id, "imported_at": imported_at,
            "source_sheet": config["worksheet"], "source_table": config["table"],
            "source_columns": len(headers), "rows_seen": rows_seen, "rows_loaded": loaded,
            "rows_rejected": rejected, "rows_deduplicated": deduplicated,
            "rejections": rejection_rows,
            "m_query_corrections": [
                "PVC thickness is corrected from the original value before replacement.",
                "Derived ColourQuantity reads the renamed original value (ColourQuant behavior).",
            ],
        }
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
        formula_book.close()
        value_book.close()


def write_reports(summary: dict[str, Any], summary_path: Path, rejection_path: Path) -> None:
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    rejection_path.parent.mkdir(parents=True, exist_ok=True)
    summary_copy = dict(summary)
    rejections = summary_copy.pop("rejections", [])
    summary_path.write_text(json.dumps(summary_copy, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with rejection_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["source_row", "product_code_raw", "severity", "rule_code", "reason"])
        writer.writeheader()
        writer.writerows(rejections)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", default=str(Path(__file__).resolve().parents[1] / "config.example.json"))
    parser.add_argument("--source", help="Override source_file from the JSON configuration")
    parser.add_argument("--database", help="Override database_path from the JSON configuration")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        config = load_config(Path(args.config))
        summary = import_masterlist(config, args.source, args.database)
        write_reports(summary, Path(config["summary_report"]), Path(config["rejection_report"]))
        printable = {key: value for key, value in summary.items() if key != "rejections"}
        print(json.dumps(printable, ensure_ascii=False, indent=2))
        return 0
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
